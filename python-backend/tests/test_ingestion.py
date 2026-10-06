import unittest
import io
from datetime import datetime, timezone
from unittest.mock import patch

from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.domain import CollectionData, DocumentData
from app.domain import ChunkData
from app.ingestion import (
    DocumentElement,
    build_chunks,
    extract_and_build_chunks,
    ingest_document,
    parse_docx_elements,
    parse_docx_text,
    parse_pdf_pages,
    parse_txt_text,
)
from app.store import MemoryStore


class IngestionTests(unittest.IsolatedAsyncioTestCase):
    def make_store_with_document(self):
        store = MemoryStore()
        document = DocumentData(
            id="document-1",
            name="notes.txt",
            checksum="checksum",
            uploaded_at=datetime.now(timezone.utc),
            status="PROCESSING",
            error=None,
        )
        collection = CollectionData(
            id="collection-1",
            name="Research",
            created_at=datetime.now(timezone.utc),
            documents={document.id: document},
        )
        store.collections[collection.id] = collection
        return store, collection, document

    def test_parse_txt_text_decodes_utf8(self):
        expected = "caf\u00e9"
        self.assertEqual(parse_txt_text(expected.encode("utf-8")), expected)

    def test_parse_txt_text_falls_back_to_utf16_and_latin1(self):
        self.assertEqual(parse_txt_text("hello".encode("utf-16")), "hello")
        self.assertEqual(parse_txt_text(b"caf\xe9"), "caf\u00e9")

    def test_parse_txt_text_decodes_utf32_before_utf16(self):
        self.assertEqual(parse_txt_text("hello".encode("utf-32")), "hello")

    def test_parse_docx_text_joins_paragraphs(self):
        document = DocxDocument()
        document.add_paragraph("First paragraph")
        document.add_paragraph("Second paragraph")
        buffer = io.BytesIO()
        document.save(buffer)

        self.assertEqual(parse_docx_text(buffer.getvalue()), "First paragraph\nSecond paragraph")

    def test_parse_docx_elements_preserves_heading_order_and_table_rows(self):
        document = DocxDocument()
        document.add_heading("Overview", level=1)
        document.add_paragraph("General introduction")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Description"
        table.add_row().cells[0].text = "Example"
        buffer = io.BytesIO()
        document.save(buffer)

        elements = parse_docx_elements(buffer.getvalue())
        chunks = build_chunks(elements, chunk_size=200, overlap=20)

        self.assertEqual([element.kind for element in elements[:3]], ["heading", "paragraph", "table_row"])
        self.assertIn("Item | Description", [element.text for element in elements])
        table_chunk = next(chunk for chunk in chunks if "Item | Description" in chunk.text)
        self.assertEqual(table_chunk.section_path, ("Overview",))
        self.assertIn("table_row", table_chunk.element_types)
        self.assertIsNotNone(table_chunk.source_start)
        self.assertTrue(any("table 1" in source for source in table_chunk.source_locations))

    def test_parse_docx_list_paragraphs_keep_list_kind(self):
        document = DocxDocument()
        document.add_paragraph("First item", style="List Bullet")
        buffer = io.BytesIO()
        document.save(buffer)

        self.assertEqual(parse_docx_elements(buffer.getvalue())[0].kind, "list_item")

    def test_parse_pdf_pages_preserves_one_based_page_numbers(self):
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.add_blank_page(width=72, height=72)
        buffer = io.BytesIO()
        writer.write(buffer)

        self.assertEqual(parse_pdf_pages(buffer.getvalue()), [(1, ""), (2, "")])
        with self.assertRaises(Exception):
            parse_pdf_pages(b"not a pdf")

    def test_build_chunks_skips_empty_pages_and_preserves_page_numbers(self):
        chunks = build_chunks(
            [(1, "   "), (2, "alpha bravo charlie delta echo foxtrot")],
            chunk_size=20,
            overlap=5,
        )

        self.assertGreater(len(chunks), 1)
        self.assertEqual([chunk.index for chunk in chunks], list(range(1, len(chunks) + 1)))
        self.assertEqual({chunk.page_number for chunk in chunks}, {2})
        self.assertTrue(all(len(chunk.vector) == 1536 for chunk in chunks))

    def test_build_chunks_embeds_all_chunks_with_the_configured_provider(self):
        batches = []

        def embed_many(texts):
            batches.append(texts)
            return [[float(index)] * 1536 for index, _ in enumerate(texts)]

        chunks = build_chunks(
            [(1, "alpha bravo charlie delta echo foxtrot golf hotel")],
            chunk_size=20,
            overlap=5,
            embed_many=embed_many,
        )

        self.assertEqual(len(batches), 1)
        self.assertEqual(batches[0], [chunk.text for chunk in chunks])
        self.assertEqual([chunk.vector[0] for chunk in chunks], list(map(float, range(len(chunks)))))

    def test_build_chunks_returns_empty_when_every_page_is_blank(self):
        self.assertEqual(build_chunks([(1, "  "), (None, "")], chunk_size=200, overlap=20), [])

    def test_build_chunks_retains_pdf_page_and_source_metadata(self):
        chunks = build_chunks(
            [DocumentElement("A paragraph", page_number=4, source_location="page 4, block 1")],
            chunk_size=200,
            overlap=20,
        )

        self.assertEqual(chunks[0].page_number, 4)
        self.assertEqual(chunks[0].source_start, chunks[0].source_end)
        self.assertEqual(chunks[0].element_types, ("paragraph",))

    def test_image_only_pdf_returns_actionable_ocr_error(self):
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        buffer = io.BytesIO()
        writer.write(buffer)

        with self.assertRaisesRegex(ValueError, "require OCR"):
            extract_and_build_chunks("pdf", buffer.getvalue(), chunk_size=200, overlap=20)

    def test_extract_and_build_chunks_rejects_unknown_format(self):
        with self.assertRaisesRegex(ValueError, "Unsupported document format"):
            extract_and_build_chunks("unknown", b"content", chunk_size=200, overlap=20)

    async def test_ingest_document_marks_document_ready_and_records_audit(self):
        store, collection, document = self.make_store_with_document()

        await ingest_document(
            store,
            store.collections,
            chunk_size=30,
            overlap=5,
            collection_id=collection.id,
            document_id=document.id,
            extension="txt",
            payload=b"A readable document with several words.",
            actor="editor",
            request_id="request-1",
        )

        self.assertEqual(document.status, "READY")
        self.assertIsNone(document.error)
        self.assertGreater(document.chunk_count, 0)
        self.assertEqual(store.audit_events[0].action, "DOCUMENT_INGESTION_COMPLETED")
        self.assertEqual(store.audit_events[0].outcome, "READY")

    async def test_ingest_document_marks_empty_text_failed_with_actionable_error(self):
        store, collection, document = self.make_store_with_document()

        await ingest_document(
            store,
            store.collections,
            chunk_size=100,
            overlap=10,
            collection_id=collection.id,
            document_id=document.id,
            extension="txt",
            payload=b"   \n\t",
            actor="editor",
            request_id="request-2",
        )

        self.assertEqual(document.status, "FAILED")
        self.assertEqual(document.error, "No readable text was found in this document")
        self.assertEqual(store.audit_events[0].action, "DOCUMENT_INGESTION_FAILED")
        self.assertEqual(store.audit_events[0].outcome, "FAILED")

    async def test_ingest_document_ignores_missing_collection_or_document(self):
        store, collection, _ = self.make_store_with_document()

        await ingest_document(
            store, store.collections, 100, 10, "missing", "document-1", "txt", b"content", "editor", ""
        )
        await ingest_document(
            store, store.collections, 100, 10, collection.id, "missing", "txt", b"content", "editor", ""
        )

        self.assertEqual(store.audit_events, [])

    async def test_ingest_document_rejects_documents_over_chunk_limit(self):
        store, collection, document = self.make_store_with_document()
        too_many_chunks = [ChunkData("chunk", index, None, []) for index in range(1801)]

        with patch("app.ingestion.build_chunks", return_value=too_many_chunks):
            await ingest_document(
                store,
                store.collections,
                chunk_size=200,
                overlap=20,
                collection_id=collection.id,
                document_id=document.id,
                extension="txt",
                payload=b"valid source text",
                actor="editor",
                request_id="request-3",
            )

        self.assertEqual(document.status, "FAILED")
        self.assertEqual(document.error, "Document is too large after extraction; split it into smaller files")

    async def test_ingest_document_hides_pdf_and_docx_parser_errors(self):
        store, collection, pdf_document = self.make_store_with_document()
        docx_document = DocumentData(
            id="document-2",
            name="broken.docx",
            checksum="checksum-2",
            uploaded_at=datetime.now(timezone.utc),
            status="PROCESSING",
            error=None,
        )
        collection.documents[docx_document.id] = docx_document

        for document, extension in ((pdf_document, "pdf"), (docx_document, "docx")):
            await ingest_document(
                store,
                store.collections,
                chunk_size=200,
                overlap=20,
                collection_id=collection.id,
                document_id=document.id,
                extension=extension,
                payload=b"not a valid document",
                actor="editor",
                request_id="parser-error",
            )
            self.assertEqual(document.status, "FAILED")
            expected_error = (
                "PDF could not be opened or is damaged"
                if extension == "pdf"
                else "Document processing failed. Check provider settings and retry."
            )
            self.assertEqual(document.error, expected_error)

        self.assertEqual(len(store.audit_events), 2)
        self.assertTrue(all(event.action == "DOCUMENT_INGESTION_FAILED" for event in store.audit_events))


if __name__ == "__main__":
    unittest.main()
